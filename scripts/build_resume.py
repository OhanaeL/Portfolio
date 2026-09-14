"""Render the résumé as .docx (editable source) and .pdf (via headless Edge) from one content definition."""

import html
import subprocess
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

REPO = Path(__file__).resolve().parents[1]
DOCX = REPO / "content/about/resume.docx"
PDF = REPO / "public/media/about/resume.pdf"
HTML = REPO / "content/about/resume.html"
EDGE = r"C:\Program Files (x86)\Microsoft\Edge\Application\msedge.exe"

# ---------------------------------------------------------------- content
# A "part" is a str, ("bold", text), or ("link", text, url).
NAME = "HTIN LINN"
ROLE = "Software Engineer"
CONTACT = [
    "Burmese | Bangkok, Thailand | +66 65 556 6035 | ",
    ("link", "htinlinn.dev.03@gmail.com", "mailto:htinlinn.dev.03@gmail.com"),
    " | ",
    ("link", "LinkedIn", "https://www.linkedin.com/in/htin-linn-b599711a1/"),
    " | ",
    ("link", "GitHub", "https://github.com/OhanaeL"),
    " | ",
    ("link", "Portfolio", "https://ohanael.github.io/Portfolio/"),
]
SUMMARY = (
    "Software engineer, Python first. I build production backend and AI services: REST and gRPC APIs, "
    "MCP servers and LLM integrations, deployed on Kubernetes with CI/CD and automated tests. Currently "
    "building the runtime that AI agents run on, in Python, Rust and Go."
)
LANGUAGES = "English (IELTS 8.0), Burmese (Native), Thai (Basic), Japanese (JLPT N4)"
SKILLS = [
    ("Programming", "Python, TypeScript/JavaScript, Rust, Go, Java, C#, HTML/CSS"),
    ("Frameworks", "FastAPI, Django, Node.js, Express, React, Next.js"),
    ("AI/ML", "LLM APIs, agent runtimes & native tool calling, LangGraph, LangChain, RAG, MCP, Gemini, OCR, Computer Vision (OpenCV, MediaPipe)"),
    ("Backend & APIs", "gRPC/Protobuf, RESTful design, async programming, SQL (PostgreSQL, MySQL), MongoDB, Redis"),
    ("DevOps & Tools", "Docker, Kubernetes, CI/CD (GitHub Actions), Git, Grafana, OpenTelemetry, pytest, Jira, AWS"),
]
EXPERIENCE = [
    {
        "title": "Software Engineer",
        "org": [("link", "General Magick Industries", "https://generalmagickindustries.com/"), " (", ("link", "MagickMind", "https://magickmind.ai/"), ")"],
        "tail": "Full-time | Apr 2026 - Present",
        "bullets": [
            [("bold", "Top committer"), " since joining on the reasoning, memory and supervisor services. Contributor to Mindroid, the open-source Rust agent runtime."],
            ["Built the agent supervisor: it mints credentials for each agent and runs it as an isolated process that joins workspaces over pub/sub. Added caching to its execution loop, cutting agent turn latency from ", ("bold", "4s to 2.5s"), "."],
            ["Refactored the reasoning service, cutting fast-path latency ", ("bold", "40%"), ". Implemented RLM, Lambda-RLM, Mixture of Judges and MCTS; on GPQA they score ", ("bold", "22–34 points"), " above a direct call with the same model."],
            ["Shipped native tool calling and mid-turn escalation from a fast model to a frontier model across the gateway, reasoning service and supervisor, plus an OpenAI-compatible chat-completions route with tool and image support."],
        ],
    },
    {
        "title": "Associate Fullstack Engineer",
        "org": [("link", "Brillar", "https://www.brillar.io/"), " (", ("link", "Atenxion", "https://www.atenxion.ai/"), ")"],
        "tail": "Singapore (Remote) | May 2025 - Apr 2026",
        "bullets": [
            ["Owned two Python microservices and built document processing & web crawling pipelines."],
            ["Built FastAPI services using LangGraph agents, RAG, MCP servers and AI-driven UI widgets."],
            ["Built the shared CRM integration service (Zoho, Chatwoot, Zendesk) used by several products."],
            ["Cut container build and startup time by ", ("bold", "30%"), " with Docker changes, and added OpenTelemetry and Grafana."],
            ["Maintained pytest suites in CI/CD for stable weekly releases and reviewed microservice integrations."],
        ],
    },
    {
        "title": "Software Developer Intern",
        "org": [("link", "Rangsit International College", "https://rsuip.org/")],
        "tail": "Thailand | Aug 2024 - Dec 2025 | Part-time from May 2025",
        "bullets": [
            ["Built a QR-based attendance system (Next.js/Express) saving ", ("bold", "15 minutes"), " per class per staff."],
            ["Led Express backend & React frontend for a job fair platform used by ", ("bold", "300+ students and 100+ companies"), "."],
            ["Built a Python/Django course recommendation system used by ", ("bold", "100+ students"), "."],
        ],
    },
    {
        "title": "Freelance Software Developer",
        "org": ["Self Employed"],
        "tail": "Myanmar | Jan 2021 - Feb 2022",
        "bullets": [
            ["Built a Python/Bootstrap online ordering system for a local restaurant that cut average order time by ", ("bold", "~3 minutes"), ", plus Shopee inventory-sync integrations for small businesses."],
        ],
    },
]
ACHIEVEMENTS = [
    [("bold", "1st Place"), ", CIMSO Hospitality ERP Hackathon (Mar 2025) – Built ", ("bold", "SPLASH Golf Club"), " (React + ERP integration) with an AI voice chatbot for bookings."],
    [("bold", "3rd Place, Hack the Zodiac Hackathon (Apr 2024)"), " – Built NextThai and EcoMart, two AI consumer apps."],
]
DEGREE = "Bachelor of Science (Information and Communication Technology)"
SCHOOL = ["Rangsit University | Pathum Thani, Thailand | 2022-2025 | ", ("bold", "GPA: 3.96/4"), " | ", ("bold", "First Class Honours")]
EDU_BULLETS = [["Mentored classmates and foreign students in programming fundamentals."]]

# ---------------------------------------------------------------- variants
# `python scripts/build_resume.py systems` writes resume-systems.{docx,pdf} to
# content/about/ (not served by the site): the same facts, tilted towards
# systems, Linux and security work for roles that are not about LLMs.
VARIANTS = {
    "systems": {
        "SUMMARY": (
            "Software engineer working in Rust, Python and Go on backend services that run in production on Linux "
            "and Kubernetes. Currently building a Rust control plane that supervises one OS process per AI agent and "
            "mints its credentials. Comfortable in unfamiliar codebases (17 services, four languages); CI/CD, tests "
            "and observability are part of finishing the job."
        ),
        "SKILLS": [
            ("Programming", "Rust, Python, Go, Bash/Shell, TypeScript, Java, C#"),
            ("Systems & platform", "Linux containers, Docker, Kubernetes, etcd service discovery, gRPC/Protobuf, Redis (clustered), PostgreSQL, MySQL, MongoDB, RabbitMQ, Centrifugo pub/sub"),
            ("Security practice", "per-service credentials, tenant isolation, input and host validation on user-supplied URLs, secrets kept out of repositories, dependency CVE remediation, SonarQube quality gates"),
            ("CI/CD & testing", "GitHub Actions, cargo test/clippy, pytest, multi-stage Docker builds, OpenTelemetry, Grafana"),
            ("AI (secondary)", "LLM APIs, agent runtimes, LangGraph, RAG, MCP"),
        ],
        "GMI_BULLETS": [
            ["Built the agent supervisor in ", ("bold", "Rust"), ": a control plane that mints per-agent credentials and runs one isolated OS process per agent, enforcing a ", ("bold", "tenant boundary"), " on every request that acts on someone's behalf."],
            ["Hardened the gateway's public chat route (Go): request-size cap, ", ("bold", "host validation"), " on user-supplied URLs before the server fetches them, and structured tool calls in place of text parsing."],
            ["Maintenance across three services: remediated ", ("bold", "vendored CVEs"), " in container images (stripped pip from the uv-managed interpreter, pinned setuptools), namespaced Redis keys per service, and made message ingest ", ("bold", "idempotent"), " so retries cannot duplicate records."],
            ["Cut agent turn latency from ", ("bold", "4s to 2.5s"), " and the reasoning service's fast path by ", ("bold", "40%"), " with targeted caching, measured before and after."],
            [("bold", "Top committer"), " since joining on three production services; every change ships with tests through CI to Kubernetes. Contributor to Mindroid, the company's open-source Rust runtime."],
        ],
        "BRILLAR_BULLETS": [
            ["Owned two Python (FastAPI) microservices end to end: design, implementation and production operations."],
            ["Cut container build and startup time by ", ("bold", "30%"), " with Docker changes; added OpenTelemetry and Grafana."],
            ["Maintained pytest suites in CI/CD for stable weekly releases; reviewed other teams' integrations against the shared CRM integration service (Zoho, Chatwoot, Zendesk)."],
            ["Built the document ingestion pipeline that turns uploads and crawled pages into text for downstream services."],
        ],
        "ROLE": "Software Engineer (Rust / Python / Go)",
    },
    # Forward-deployed / solutions roles: the whole stack, reuse of building
    # blocks, security and compliance practice named, documentation.
    "fde": {
        "SUMMARY": (
            "Software engineer shipping production code across backend, frontend, data and ops, in Rust, Go, Python "
            "and TypeScript. I build and operate the runtime AI agents run on: an agent supervisor, a reasoning service "
            "and agent memory, on Kubernetes with CI/CD, tests and observability. Comfortable in unfamiliar codebases "
            "(17 services, four languages), and I write things down so the next team can reuse them."
        ),
        "SKILLS": [
            ("Programming", "Python, Rust, Go, TypeScript/JavaScript, Java, C#, Bash"),
            ("AI stack", "LLM APIs (OpenAI, Gemini, Claude), agent runtimes, tool calling, LangGraph, RAG, MCP"),
            ("Backend & data", "FastAPI, gRPC/Protobuf, REST, SQL (PostgreSQL, MySQL), MongoDB, Redis, RabbitMQ"),
            ("Frontend", "React, Next.js, Tailwind, TypeScript"),
            ("Ops & security", "Docker, Kubernetes, GitHub Actions, OpenTelemetry, Grafana, pytest, SonarQube gates; per-service credentials, tenant isolation, host validation, CVE remediation"),
        ],
        "GMI_BULLETS": [
            ["Built the agent supervisor in Rust: it mints per-agent credentials, runs one isolated process per agent and enforces a ", ("bold", "tenant boundary"), " on every request. Agents join workspaces as participants, so several share a room and every message is attributed."],
            ["Shipped native tool calling and mid-turn escalation from a fast model to a frontier model across the gateway (Go), reasoning service (Python) and runtime (Rust), through an OpenAI-compatible route with request-size caps and host validation on user-supplied URLs."],
            ["Reused and extended building blocks: Mindroid, the company's open-source Rust agent SDK, is the generic runtime and the supervisor the product layer; generic stages and tools go back into the SDK."],
            ["Implemented and benchmarked reasoning strategies (RLM, Lambda-RLM, Mixture of Judges, MCTS): ", ("bold", "22–34 points"), " above a direct call on GPQA with the same model. Cut agent turn latency from ", ("bold", "4s to 2.5s"), " and fast-path latency ", ("bold", "40%"), "."],
            [("bold", "Top committer"), " since joining on three production services; wrote the case studies and diagrams that explain the runtime to non-engineers."],
        ],
        "BRILLAR_BULLETS": [
            ["Owned two Python (FastAPI) microservices end to end on a remote team, from design to production operations."],
            ["Built the shared CRM integration service (Zoho, Chatwoot, Zendesk) used by several products, and the document ingestion pipeline that turns uploads and crawled pages into text for LLM workflows."],
            ["Built FastAPI services using LangGraph agents, RAG, MCP servers and AI-driven UI widgets; cut container build and startup time by ", ("bold", "30%"), "; added OpenTelemetry and Grafana."],
        ],
        "ROLE": "Software Engineer",
        "FREELANCE_BULLETS": [["Built a Python/Bootstrap restaurant ordering system (order time down ", ("bold", "~3 minutes"), ") and Shopee inventory sync."]],
    },
}


def apply_variant(name):
    global SUMMARY, SKILLS, ROLE, DOCX, PDF, HTML
    v = VARIANTS[name]
    SUMMARY, SKILLS, ROLE = v["SUMMARY"], v["SKILLS"], v["ROLE"]
    EXPERIENCE[0]["bullets"] = v["GMI_BULLETS"]
    EXPERIENCE[1]["bullets"] = v["BRILLAR_BULLETS"]
    if "FREELANCE_BULLETS" in v:
        EXPERIENCE[3]["bullets"] = v["FREELANCE_BULLETS"]
    DOCX = REPO / f"content/about/resume-{name}.docx"
    PDF = REPO / f"content/about/resume-{name}.pdf"
    HTML = REPO / f"content/about/resume-{name}.html"


# ---------------------------------------------------------------- docx
BLUE = RGBColor(0x2E, 0x74, 0xB5)
FONT = "Times New Roman"


def _run(p, text, bold=False, italic=False, size=None, color=None):
    r = p.add_run(text)
    r.bold, r.italic = bold, italic
    r.font.name = FONT
    r._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    if size:
        r.font.size = Pt(size)
    if color:
        r.font.color.rgb = color


def _link(p, text, url, italic=False):
    rid = p.part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    h = OxmlElement("w:hyperlink")
    h.set(qn("r:id"), rid)
    r = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    fonts = OxmlElement("w:rFonts")
    for a in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        fonts.set(qn(a), FONT)
    rpr.append(fonts)
    if italic:
        rpr.append(OxmlElement("w:i"))
    c = OxmlElement("w:color")
    c.set(qn("w:val"), "0563C1")
    rpr.append(c)
    u = OxmlElement("w:u")
    u.set(qn("w:val"), "single")
    rpr.append(u)
    r.append(rpr)
    t = OxmlElement("w:t")
    t.text = text
    t.set(qn("xml:space"), "preserve")
    r.append(t)
    h.append(r)
    p._p.append(h)


def _parts(p, parts, italic=False):
    for part in parts:
        if isinstance(part, str):
            _run(p, part, italic=italic)
        elif part[0] == "bold":
            _run(p, part[1], bold=True, italic=italic)
        else:
            _link(p, part[1], part[2], italic=italic)


def _para(doc, before=0, after=2):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    return p


def _heading(doc, text):
    _run(_para(doc, before=8, after=2), text, bold=True, size=12, color=BLUE)


def _bullet(doc, parts):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(1)
    _parts(p, parts)


def _labelled(doc, label, text):
    p = _para(doc, after=1)
    _run(p, label + ": ", bold=True)
    _run(p, text)


def build_docx():
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
    sec.top_margin = sec.bottom_margin = Cm(1.5)
    sec.left_margin = sec.right_margin = Cm(1.8)
    normal = doc.styles["Normal"]
    normal.font.name = FONT
    normal.font.size = Pt(10)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)

    _run(_para(doc, after=0), NAME, bold=True, size=24)
    _run(_para(doc, after=2), ROLE, size=15)
    _parts(_para(doc, after=0), CONTACT)

    _heading(doc, "SUMMARY")
    _run(_para(doc, after=1), SUMMARY)
    _labelled(doc, "Languages", LANGUAGES)

    _heading(doc, "EXPERIENCE")
    for job in EXPERIENCE:
        p = _para(doc, before=6, after=2)
        _run(p, job["title"], bold=True)
        _run(p, " | ")
        _parts(p, job["org"], italic=True)
        _run(p, " | " + job["tail"], italic=True)
        for b in job["bullets"]:
            _bullet(doc, b)

    _heading(doc, "SKILLS")
    for label, text in SKILLS:
        _labelled(doc, label, text)

    _heading(doc, "PROJECTS & ACHIEVEMENTS")
    for b in ACHIEVEMENTS:
        _bullet(doc, b)

    _heading(doc, "EDUCATION")
    _run(_para(doc, after=0), DEGREE, bold=True)
    _parts(_para(doc, after=1), SCHOOL, italic=True)
    for b in EDU_BULLETS:
        _bullet(doc, b)

    doc.core_properties.title = f"{NAME.title()} | {ROLE}"
    doc.core_properties.author = NAME.title()
    doc.save(DOCX)


# ---------------------------------------------------------------- html -> pdf
CSS = """
@page { size: A4; margin: 13mm 16mm; }
body { font-family: "Times New Roman", Times, serif; font-size: 10pt; line-height: 1.25; color: #000; margin: 0; }
a { color: #0563C1; text-decoration: underline; }
.name { font-size: 24pt; font-weight: bold; line-height: 1.05; }
.role { font-size: 15pt; margin: 2pt 0 4pt; }
.contact { font-size: 10pt; }
h2 { color: #2E74B5; font-size: 12pt; margin: 8pt 0 2pt; }
p { margin: 0 0 1.5pt; }
ul { margin: 0 0 1pt; padding-left: 16pt; }
li { margin: 0 0 1pt; }
.job { margin-top: 5pt; margin-bottom: 1pt; }
.job em, .school em { font-style: italic; }
"""


def _h(parts, italic=False):
    out = []
    for part in parts:
        if isinstance(part, str):
            out.append(html.escape(part))
        elif part[0] == "bold":
            out.append(f"<b>{html.escape(part[1])}</b>")
        else:
            out.append(f'<a href="{html.escape(part[2])}">{html.escape(part[1])}</a>')
    s = "".join(out)
    return f"<em>{s}</em>" if italic else s


def _ul(items):
    return "<ul>" + "".join(f"<li>{_h(b)}</li>" for b in items) + "</ul>"


def build_html():
    exp = "".join(
        f'<p class="job"><b>{html.escape(j["title"])}</b> | {_h(j["org"], italic=True)} | <em>{html.escape(j["tail"])}</em></p>'
        + _ul(j["bullets"])
        for j in EXPERIENCE
    )
    skills = "".join(f"<p><b>{html.escape(k)}:</b> {html.escape(v)}</p>" for k, v in SKILLS)
    doc = f"""<!doctype html><html><head><meta charset="utf-8"><title>{NAME.title()} | {ROLE}</title><style>{CSS}</style></head><body>
<div class="name">{NAME}</div><div class="role">{ROLE}</div><div class="contact">{_h(CONTACT)}</div>
<h2>SUMMARY</h2><p>{html.escape(SUMMARY)}</p><p><b>Languages:</b> {html.escape(LANGUAGES)}</p>
<h2>EXPERIENCE</h2>{exp}
<h2>SKILLS</h2>{skills}
<h2>PROJECTS &amp; ACHIEVEMENTS</h2>{_ul(ACHIEVEMENTS)}
<h2>EDUCATION</h2><p><b>{html.escape(DEGREE)}</b></p><p class="school">{_h(SCHOOL, italic=True)}</p>{_ul(EDU_BULLETS)}
</body></html>"""
    HTML.write_text(doc, encoding="utf-8")


def build_pdf():
    subprocess.run(
        [EDGE, "--headless=new", "--disable-gpu", "--no-pdf-header-footer", f"--print-to-pdf={PDF}", HTML.as_uri()],
        check=True, timeout=120, capture_output=True,
    )


if __name__ == "__main__":
    import sys

    if len(sys.argv) > 1:
        apply_variant(sys.argv[1])
    build_docx()
    build_html()
    build_pdf()
    HTML.unlink(missing_ok=True)
    print("docx", DOCX.stat().st_size, "pdf", PDF.stat().st_size)
